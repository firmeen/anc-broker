#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read-only, consistent MySQL snapshot of the four ANC tables into four Excel sheets."""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
import time
import zipfile
from datetime import date, datetime, time as date_time
from decimal import Decimal
from pathlib import Path

import pymysql
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from rich.console import Console, Group
from rich.live import Live
from rich.panel import Panel
from rich.table import Table
from rich.text import Text


TABLE_SHEETS = (
    ("anc_competitor_insurance_v4", "anc_competitor_insurance_v4"),
    ("anc_competitor_target_v4", "anc_competitor_target_v4"),
    ("anc_competitor_branch_v4", "anc_competitor_branch_v4"),
    ("anc_competitor_package_checkpoint_v4", "anc_package_checkpoint_v4"),
)
DEFAULT_OUTPUT = "anc_competitor_db_snapshot.xlsx"
MAX_EXCEL_ROWS = 1_048_576
MAX_DATA_ROWS = MAX_EXCEL_ROWS - 1
MAX_TEXT = 32_767
BATCH_SIZE = 1000
console = Console()


class ExportError(Exception):
    """Clear, actionable failure that must not produce a completed snapshot."""


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="Read-only ANC MySQL -> four-sheet Excel snapshot")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Output .xlsx filename")
    parser.add_argument("--overwrite", action="store_true", help="Replace existing output after validation")
    args = parser.parse_args(argv)
    if Path(args.output).suffix.lower() != ".xlsx":
        parser.error("--output must end with .xlsx")
    return args


def setting(*names, default=None):
    for name in names:
        value = os.getenv(name)
        if value not in (None, ""):
            return value
    return default


def connect_mysql():
    load_dotenv()
    user = setting("DB_USER", "MYSQL_USER")
    if not user:
        raise ExportError("DB_USER (or MYSQL_USER) is required in .env")
    port = int(setting("DB_PORT", "MYSQL_PORT", default="3307"))
    name = setting("DB_NAME", "MYSQL_DATABASE", default="dbd")
    connection = pymysql.connect(
        host=setting("DB_HOST", "MYSQL_HOST", default="127.0.0.1"),
        port=port,
        user=user,
        password=setting("DB_PASSWORD", "MYSQL_PASSWORD", default=""),
        database=name,
        charset="utf8mb4",
        autocommit=True,
        connect_timeout=15,
        read_timeout=120,
        cursorclass=pymysql.cursors.Cursor,
    )
    return connection, name, port


def check_tables(connection):
    """Require all four source tables to be InnoDB before claiming MVCC consistency."""
    names = [table for table, _ in TABLE_SHEETS]
    markers = ",".join(["%s"] * len(names))
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT TABLE_NAME, ENGINE FROM information_schema.TABLES "
            f"WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME IN ({markers})",
            names,
        )
        found = {name: engine for name, engine in cursor.fetchall()}
    for table in names:
        if table not in found:
            raise ExportError(f"Required MySQL table missing: {table}")
        if found[table].upper() != "INNODB":
            raise ExportError(f"Consistent snapshot requires InnoDB: {table} uses {found[table]}")


def start_snapshot(connection):
    with connection.cursor() as cursor:
        cursor.execute("SET SESSION TRANSACTION ISOLATION LEVEL REPEATABLE READ")
        cursor.execute("START TRANSACTION WITH CONSISTENT SNAPSHOT, READ ONLY")


def count_tables(connection):
    counts = {}
    with connection.cursor() as cursor:
        for table, _ in TABLE_SHEETS:
            cursor.execute(f"SELECT COUNT(*) FROM `{table}`")
            count = int(cursor.fetchone()[0])
            if count > MAX_DATA_ROWS:
                raise ExportError(
                    f"{table}: {count:,} rows exceed Excel limit {MAX_DATA_ROWS:,} data rows/sheet"
                )
            counts[table] = count
    return counts


def safe_xml_text(value, table, row_number, column):
    if len(value) > MAX_TEXT:
        raise ExportError(
            f"{table} row {row_number} column {column}: {len(value):,} characters "
            f"exceed Excel's {MAX_TEXT:,}-character cell limit"
        )
    for c in value:
        point = ord(c)
        if not (point in (9, 10, 13) or 32 <= point <= 0xD7FF
                or 0xE000 <= point <= 0xFFFD or 0x10000 <= point <= 0x10FFFF):
            raise ExportError(
                f"{table} row {row_number} column {column}: "
                f"character U+{point:04X} cannot be stored in Excel XML"
            )
    return value


def excel_cell(sheet, value, table, row_number, column):
    if value is None:
        return None  # SQL NULL becomes a blank cell (not distinguishable from empty text in Excel).
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        text = safe_xml_text(value, table, row_number, column)
        cell = WriteOnlyCell(sheet, value=text)
        cell.data_type = "s"  # Force text, including values beginning with =, +, -, @.
        return cell
    if isinstance(value, int):
        if len(str(abs(value))) > 15:
            cell = WriteOnlyCell(sheet, value=str(value))
            cell.data_type = "s"
            return cell
        return value
    if isinstance(value, Decimal):
        cell = WriteOnlyCell(sheet, value=format(value, "f"))
        cell.data_type = "s"  # Avoid binary-floating-point corruption of exact MySQL DECIMAL.
        return cell
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            raise ExportError(f"{table} row {row_number} column {column}: timezone-aware datetime unsupported")
        if value.year < 1900 or value.microsecond:
            return excel_cell(sheet, value.isoformat(sep=" "), table, row_number, column)
        return value
    if isinstance(value, date):
        if value.year < 1900:
            return excel_cell(sheet, value.isoformat(), table, row_number, column)
        return value
    if isinstance(value, date_time):
        return excel_cell(sheet, value.isoformat(), table, row_number, column)
    raise ExportError(
        f"{table} row {row_number} column {column}: unsupported data type {type(value).__name__}"
    )


class ExportDashboard:
    """Small Rich display; counters reflect successfully appended Excel rows only."""

    def __init__(self, database="--", port="--", output="--"):
        self.database, self.port, self.output = database, port, output
        self.status = "CONNECTING"
        self.stage = "Connecting to MySQL"
        self.started = time.monotonic()
        self.current_started = self.started
        self.current = None
        self.totals = {name: None for name, _ in TABLE_SHEETS}
        self.written = {name: 0 for name, _ in TABLE_SHEETS}
        self.done = set()
        self.live = None
        self.last_refresh = 0.0

    def render(self):
        top = Table.grid(expand=True)
        top.add_column(ratio=1)
        top.add_column(justify="right")
        top.add_row(Text("ANC DATABASE EXPORTER  |  MYSQL → EXCEL", style="bold cyan"),
                    Text(self.status, style="bold green" if self.status == "COMPLETED" else "yellow"))
        top.add_row(f"Database: {self.database}  |  Port: {self.port}",
                    f"Elapsed: {time.monotonic() - self.started:.1f}s")
        top.add_row("Output: " + str(self.output), "")
        table = Table(expand=True, box=None, show_lines=False, pad_edge=False)
        table.add_column("Table", ratio=3, overflow="ellipsis", no_wrap=True)
        table.add_column("Written / Total", justify="right", ratio=2)
        table.add_column("Status", justify="right", ratio=1)
        for full, _ in TABLE_SHEETS:
            label = full.removeprefix("anc_competitor_")
            total = self.totals[full]
            state = "DONE" if full in self.done else "EXPORTING" if full == self.current else "PENDING"
            table.add_row(label, f"{self.written[full]:,} / {total:,}" if total is not None else "--", state)
        complete = len(self.done)
        current_rows = self.written.get(self.current, 0)
        current_total = self.totals.get(self.current)
        elapsed = time.monotonic() - self.current_started
        rate = current_rows / elapsed if self.current is not None and elapsed > 0 else 0
        fraction = current_rows / current_total if current_total else (1.0 if self.current in self.done else 0.0)
        filled = min(30, max(0, int(30 * fraction)))
        eta = (current_total - current_rows) / rate if current_total is not None and rate > 0 and elapsed >= 1 else None
        eta_label = f"{eta:.0f}s" if eta is not None else "--"
        detail = (
            f"Tables: {complete}/4  |  Stage: {self.stage}\n"
            f"Current: {self.current or '--'}\n"
            f"[cyan]{'━' * filled}[/cyan][dim]{'─' * (30 - filled)}[/dim] {fraction * 100:.1f}%\n"
            f"Rows: {current_rows:,}  |  Rate: {rate:,.0f} rows/s  |  "
            f"Elapsed: {elapsed:.1f}s  |  ETA: {eta_label}"
        )
        return Panel(Group(top, table, Panel(detail, title="PROGRESS", border_style="cyan")),
                     title="READ-ONLY DATABASE SNAPSHOT", border_style="cyan")

    def start(self):
        self.live = Live(self.render(), console=console, refresh_per_second=4, transient=True)
        self.live.start()

    def update(self, force=False):
        now = time.monotonic()
        if self.live and (force or now - self.last_refresh >= 0.25):
            self.last_refresh = now
            self.live.update(self.render(), refresh=True)

    def close(self):
        if self.live:
            self.live.stop()
            self.live = None


def export_table(connection, workbook, table, worksheet, expected, dashboard):
    dashboard.status, dashboard.stage = "EXPORTING", "Reading and writing rows"
    dashboard.current, dashboard.current_started = table, time.monotonic()
    dashboard.update(force=True)
    sheet = workbook.create_sheet(title=worksheet)
    written = 0
    # SSCursor reads the server result incrementally rather than buffering the full table.
    with connection.cursor(pymysql.cursors.SSCursor) as cursor:
        cursor.execute(f"SELECT * FROM `{table}`")
        columns = [entry[0] for entry in cursor.description]
        if not columns:
            raise ExportError(f"{table}: missing result columns")
        header = []
        for name in columns:
            cell = WriteOnlyCell(sheet, value=name)
            cell.data_type = "s"
            cell.font = Font(bold=True)
            header.append(cell)
        sheet.append(header)
        # Stable order is not guaranteed without ORDER BY; contents and multiplicity are preserved.
        while True:
            batch = cursor.fetchmany(BATCH_SIZE)
            if not batch:
                break
            for record in batch:
                row_number = written + 2
                cells = [excel_cell(sheet, value, table, row_number, column)
                         for column, value in zip(columns, record)]
                sheet.append(cells)
                written += 1
                dashboard.written[table] = written
            dashboard.update()
    if written != expected:
        raise ExportError(f"{table}: expected {expected:,} rows, wrote {written:,}; snapshot incomplete")
    dashboard.done.add(table)
    dashboard.update(force=True)


def validate_workbook(path, counts):
    # Validate file structure and ZIP CRC before final installation.
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ExportError("Generated Excel archive failed ZIP integrity validation")
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        expected_sheets = [sheet for _, sheet in TABLE_SHEETS]
        if book.sheetnames != expected_sheets:
            raise ExportError("Workbook worksheet names/order do not match the four-source contract")
        for table, worksheet in TABLE_SHEETS:
            sheet = book[worksheet]
            # For write-only sheets max_row may be unspecified; write counts were checked during streaming.
            if sheet.max_row is not None and sheet.max_row != counts[table] + 1:
                raise ExportError(f"{table}: workbook rows differ from source snapshot")
    finally:
        book.close()


def export_snapshot(connection, output, overwrite, dashboard):
    if output.exists() and not overwrite:
        raise ExportError(f"Output already exists: {output}. Use --overwrite to replace it")
    dashboard.status, dashboard.stage = "VALIDATING", "Checking source table engines"
    dashboard.update(force=True)
    check_tables(connection)
    start_snapshot(connection)
    temp_path = None
    workbook = None
    try:
        dashboard.status, dashboard.stage = "COUNTING", "Counting all four tables in one snapshot"
        dashboard.update(force=True)
        counts = count_tables(connection)
        dashboard.totals.update(counts)
        dashboard.update(force=True)
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(prefix=".anc_export_", suffix=".xlsx", dir=output.parent,
                                         delete=False) as stream:
            temp_path = Path(stream.name)
        workbook = Workbook(write_only=True)
        for table, worksheet in TABLE_SHEETS:
            export_table(connection, workbook, table, worksheet, counts[table], dashboard)
        # End the read-only consistent snapshot before potentially slow XLSX finalization.
        connection.rollback()
        dashboard.current = None
        dashboard.status, dashboard.stage = "SAVING", "Finalizing and checking Excel archive"
        dashboard.update(force=True)
        workbook.save(temp_path)
        workbook.close()
        workbook = None
        validate_workbook(temp_path, counts)
        if overwrite:
            os.replace(temp_path, output)
        else:
            # Atomic, no-clobber creation on filesystems supporting hard links.
            os.link(temp_path, output)
            temp_path.unlink()
        temp_path = None
        dashboard.status, dashboard.stage = "COMPLETED", "Workbook finalized"
        dashboard.update(force=True)
        return counts
    finally:
        if workbook is not None:
            # Complete any open write-only XML generators before discarding the workbook.
            for worksheet in workbook.worksheets:
                try:
                    worksheet.close()
                except Exception:
                    pass
            workbook.close()
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
        # Closing a snapshot with rollback also releases the read view if an error occurred.
        connection.rollback()


def main(argv=None):
    args = parse_args(argv)
    output = Path(args.output).expanduser().resolve()
    dashboard = ExportDashboard(output=output)
    connection = None
    stage = "initialization"
    dashboard.start()
    try:
        if output.exists() and not args.overwrite:
            raise ExportError(f"Output already exists: {output}. Use --overwrite")
        connection, name, port = connect_mysql()
        dashboard.database, dashboard.port = name, port
        dashboard.update(force=True)
        stage = "export"
        counts = export_snapshot(connection, output, args.overwrite, dashboard)
        dashboard.close()
        summary = Table(title="ANC DATABASE EXPORT COMPLETED", show_header=True)
        summary.add_column("Table")
        summary.add_column("Rows", justify="right")
        for table, _ in TABLE_SHEETS:
            summary.add_row(table, f"{counts[table]:,}")
        summary.add_row("TOTAL", f"{sum(counts.values()):,}")
        console.print(summary)
        console.print(f"Worksheets: 4 | File: {output} | Elapsed: {time.monotonic()-dashboard.started:.1f}s")
        return 0
    except KeyboardInterrupt:
        dashboard.status, dashboard.stage = "INTERRUPTED", "Cancelled by user"
        dashboard.close()
        console.print("[yellow]INTERRUPTED: incomplete temporary output removed; existing output preserved.[/yellow]")
        return 130
    except Exception as exc:
        dashboard.status, dashboard.stage = "FAILED", stage
        table = dashboard.current or "--"
        dashboard.close()
        console.print(f"[red]FAILED[/red] | Stage: {stage} | Table: {table} | {exc}")
        console.print("Incomplete temporary output removed; any existing output preserved.")
        return 1
    finally:
        if connection is not None:
            connection.close()
        dashboard.close()


if __name__ == "__main__":
    sys.exit(main())